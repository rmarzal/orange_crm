"""Helpers para editar el pptx preservando estilo + verificador de ajuste de texto."""
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from PIL import ImageFont
import os

# Fuente métrica-compatible con Arial/Helvetica para estimar anchos
_FONT_REG = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
_FONT_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
_cache = {}
def _font(size_px, bold=False):
    key=(round(size_px),bold)
    if key not in _cache:
        _cache[key]=ImageFont.truetype(_FONT_BOLD if bold else _FONT_REG, round(size_px))
    return _cache[key]

def text_width_px(text, size_pt, bold=False):
    f=_font(size_pt*96/72.0, bold)
    return f.getlength(text)

def wrap_lines(text, size_pt, avail_px, bold=False):
    """Cuenta líneas al envolver 'text' en avail_px."""
    lines=0
    for para in text.split("\n"):
        words=para.split(" ")
        cur=""
        if not words: lines+=1; continue
        for w in words:
            trial=(cur+" "+w).strip()
            if text_width_px(trial,size_pt,bold)<=avail_px or not cur:
                cur=trial
            else:
                lines+=1; cur=w
        lines+=1
    return lines

def check_fit(shape, size_pt, bold=False, l_margin=0.1, r_margin=0.1, t_margin=0.05, b_margin=0.05, line_factor=1.2, label=""):
    """Estima si el texto cabe en el shape. Devuelve (ok, lines, needed_in, avail_in)."""
    W=Emu(shape.width).inches; H=Emu(shape.height).inches
    avail_px=(W - l_margin - r_margin)*96
    text=shape.text_frame.text
    lines=wrap_lines(text, size_pt, avail_px, bold)
    needed_in = lines*size_pt*line_factor/72.0 + t_margin + b_margin
    avail_in = H
    ok = needed_in <= avail_in + 0.02
    if label:
        flag="OK " if ok else "OVER"
        print(f"  [{flag}] {label}: {lines} líneas, necesita {needed_in:.2f}in / caja {avail_in:.2f}in")
    return ok, lines, needed_in, avail_in

# ---- edición ----
def clear_runs(p):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)

def set_para(p, segments, size=None, color=None, font="Helvetica", bold=None):
    """segments: str o lista de (texto,bold) o texto. Reemplaza runs preservando pPr."""
    clear_runs(p)
    if isinstance(segments,str):
        segments=[(segments,bold)]
    for seg in segments:
        if isinstance(seg,str): t,b=seg,bold
        else: t,b=seg
        run=p.add_run(); run.text=t
        if size is not None: run.font.size=Pt(size)
        if b is not None: run.font.bold=b
        if font: run.font.name=font
        if color is not None: run.font.color.rgb=RGBColor.from_string(color)

def trim_paragraphs(tf, keep):
    ps=tf.paragraphs
    for p in ps[keep:]:
        p._p.getparent().remove(p._p)

def add_paragraph(tf):
    from pptx.oxml.ns import qn
    p=tf._txBody.makeelement(qn('a:p'),{})
    tf._txBody.append(p)
    from pptx.text.text import _Paragraph
    return _Paragraph(p, tf._txBody)
